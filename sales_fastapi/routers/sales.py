"""Sales analytics and export API.

All figures come from the synced ERP tables. Empty data returns explicit zeros
and ``data_available=false`` rather than fabricated metrics.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import case, func
from sqlalchemy.orm import Session

from ..database import get_db
from ..erp_sync import latest_report_date
from ..governance.audit import log_event
from ..models import ErpDispatch, ErpOrder, User
from ..security import get_current_user

router = APIRouter(prefix="/sales", tags=["sales"])


def _day(db: Session, day: str | None = None) -> str:
    return day or latest_report_date(db)


def _orders(db: Session, day: str):
    return db.query(ErpOrder).filter(ErpOrder.report_date == day).all()


def _company_rows(db: Session, day: str) -> list[dict]:
    order_map: dict[str, dict] = {}
    for customer, count, value in (
        db.query(
            ErpOrder.customer,
            func.count(ErpOrder.id),
            func.coalesce(func.sum(ErpOrder.total_value), 0),
        )
        .filter(ErpOrder.report_date == day)
        .group_by(ErpOrder.customer)
        .all()
    ):
        row = order_map.setdefault(customer or "(unknown)", {"company": customer or "(unknown)"})
        row.update({"orders": int(count or 0), "orders_value": float(value or 0)})

    dispatch_map: dict[str, dict] = {}
    for customer, scheduled, dispatched, pending in (
        db.query(
            ErpDispatch.customer,
            func.coalesce(func.sum(ErpDispatch.schedule_qty), 0),
            func.coalesce(func.sum(ErpDispatch.dispatched_qty), 0),
            func.sum(case((ErpDispatch.status != "Completed", 1), else_=0)),
        )
        .filter(ErpDispatch.report_date == day)
        .group_by(ErpDispatch.customer)
        .all()
    ):
        row = dispatch_map.setdefault(customer or "(unknown)", {"company": customer or "(unknown)"})
        row.update(
            {
                "scheduled_qty": float(scheduled or 0),
                "dispatched_qty": float(dispatched or 0),
                "pending": int(pending or 0),
            }
        )

    merged: dict[str, dict] = {}
    for source in (order_map, dispatch_map):
        for company, values in source.items():
            row = merged.setdefault(
                company,
                {
                    "company": company,
                    "orders": 0,
                    "orders_value": 0.0,
                    "scheduled_qty": 0.0,
                    "dispatched_qty": 0.0,
                    "pending": 0,
                },
            )
            row.update(values)
    for row in merged.values():
        scheduled = row.get("scheduled_qty", 0)
        row["completion_pct"] = (
            round(row.get("dispatched_qty", 0) / scheduled * 100, 1) if scheduled else 0
        )
    return sorted(
        merged.values(),
        key=lambda r: (r.get("orders_value", 0), r.get("dispatched_qty", 0)),
        reverse=True,
    )


def _daily_rows(db: Session, day: str, days: int) -> list[dict]:
    start = None
    if days > 0:
        start = (datetime.now(timezone.utc).date()).toordinal() - days + 1
    buckets: dict[str, dict] = {}
    for order_date, value in (
        db.query(ErpOrder.order_date, ErpOrder.total_value)
        .filter(ErpOrder.report_date == day)
        .all()
    ):
        key = (order_date or day or "")[:10]
        if not key:
            continue
        if start is not None:
            try:
                if datetime.fromisoformat(key).date().toordinal() < start:
                    continue
            except ValueError:
                pass
        bucket = buckets.setdefault(key, {"date": key, "orders": 0, "orders_value": 0.0})
        bucket["orders"] += 1
        bucket["orders_value"] += float(value or 0)
    return sorted(buckets.values(), key=lambda r: r["date"])


def _monthly_rows(db: Session, day: str) -> list[dict]:
    buckets: dict[str, dict] = {}
    for order_date, report_date, value in (
        db.query(ErpOrder.order_date, ErpOrder.report_date, ErpOrder.total_value)
        .filter(ErpOrder.report_date == day)
        .all()
    ):
        key = (order_date or report_date or "")[:7]
        if not key:
            continue
        bucket = buckets.setdefault(key, {"period": key, "orders": 0, "orders_value": 0.0})
        bucket["orders"] += 1
        bucket["orders_value"] += float(value or 0)
    return sorted(buckets.values(), key=lambda r: r["period"])


def _delivery(db: Session, day: str) -> dict:
    by_status = {
        status: int(count)
        for status, count in db.query(ErpDispatch.status, func.count(ErpDispatch.id))
        .filter(ErpDispatch.report_date == day)
        .group_by(ErpDispatch.status)
        .all()
    }
    plants = []
    for plant, records, scheduled, dispatched in (
        db.query(
            ErpDispatch.plant,
            func.count(ErpDispatch.id),
            func.coalesce(func.sum(ErpDispatch.schedule_qty), 0),
            func.coalesce(func.sum(ErpDispatch.dispatched_qty), 0),
        )
        .filter(ErpDispatch.report_date == day)
        .group_by(ErpDispatch.plant)
        .all()
    ):
        plants.append(
            {
                "plant": plant or "(unknown)",
                "records": int(records or 0),
                "scheduled_qty": float(scheduled or 0),
                "dispatched_qty": float(dispatched or 0),
                "completion_pct": round(float(dispatched or 0) / float(scheduled) * 100, 1)
                if scheduled
                else 0,
            }
        )
    pending = (
        db.query(ErpDispatch)
        .filter(ErpDispatch.report_date == day, ErpDispatch.status != "Completed")
        .order_by(ErpDispatch.balance_qty.desc())
        .limit(20)
        .all()
    )
    return {
        "by_status": by_status,
        "plants": sorted(plants, key=lambda r: r["dispatched_qty"], reverse=True),
        "pending": [
            {
                "dispatch_no": r.dispatch_no,
                "customer": r.customer,
                "plant": r.plant,
                "balance_qty": r.balance_qty,
                "pct_complete": r.pct_complete,
                "status": r.status,
            }
            for r in pending
        ],
    }


@router.get("/overview")
def overview(
    day: str | None = None, user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    report_date = _day(db, day)
    orders = _orders(db, report_date)
    dispatch = _delivery(db, report_date)
    order_value = sum(float(r.total_value or 0) for r in orders)
    scheduled = sum(
        float(r.schedule_qty or 0)
        for r in db.query(ErpDispatch).filter(ErpDispatch.report_date == report_date).all()
    )
    dispatched = sum(
        float(r.dispatched_qty or 0)
        for r in db.query(ErpDispatch).filter(ErpDispatch.report_date == report_date).all()
    )
    companies = _company_rows(db, report_date)
    return {
        "report_date": report_date,
        "data_available": bool(report_date),
        "totals": {
            "orders": len(orders),
            "orders_value": round(order_value, 2),
            "average_order_value": round(order_value / len(orders), 2) if orders else 0,
            "scheduled_qty": round(scheduled, 2),
            "dispatched_qty": round(dispatched, 2),
            "completion_pct": round(dispatched / scheduled * 100, 1) if scheduled else 0,
            "customers": len(companies),
            "pending_dispatch": sum(
                1
                for r in db.query(ErpDispatch)
                .filter(ErpDispatch.report_date == report_date, ErpDispatch.status != "Completed")
                .all()
            )
            if report_date
            else 0,
        },
        "delivery": dispatch,
        "top_companies": companies[:8],
        "monthly": _monthly_rows(db, report_date),
    }


@router.get("/companies")
def companies(
    day: str | None = None,
    q: str = "",
    limit: int = Query(default=100, ge=1, le=1000),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = _company_rows(db, _day(db, day))
    if q:
        needle = q.lower()
        rows = [r for r in rows if needle in r["company"].lower()]
    return rows[:limit]


@router.get("/daily")
def daily(
    days: int = Query(default=30, ge=1, le=365),
    day: str | None = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return _daily_rows(db, _day(db, day), days)


@router.get("/monthly")
def monthly(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return _monthly_rows(db, _day(db))


@router.get("/delivery")
def delivery(
    day: str | None = None, user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    return _delivery(db, _day(db, day))


def _export_rows(db: Session, dataset: str) -> tuple[list[str], list[list]]:
    day = _day(db)
    if dataset == "companies":
        rows = _company_rows(db, day)
        return (
            [
                "Company",
                "Orders",
                "Orders value",
                "Scheduled qty",
                "Dispatched qty",
                "Completion %",
                "Pending",
            ],
            [
                [
                    r["company"],
                    r["orders"],
                    r["orders_value"],
                    r["scheduled_qty"],
                    r["dispatched_qty"],
                    r["completion_pct"],
                    r["pending"],
                ]
                for r in rows
            ],
        )
    if dataset == "daily":
        rows = _daily_rows(db, day, 90)
        return ["Date", "Orders", "Orders value"], [
            [r["date"], r["orders"], r["orders_value"]] for r in rows
        ]
    if dataset == "dispatch":
        rows = db.query(ErpDispatch).filter(ErpDispatch.report_date == day).all()
        return (
            [
                "Dispatch",
                "Customer",
                "Plant",
                "Sales person",
                "Scheduled",
                "Dispatched",
                "Balance",
                "Completion %",
                "Status",
            ],
            [
                [
                    r.dispatch_no,
                    r.customer,
                    r.plant,
                    r.sales_person,
                    r.schedule_qty,
                    r.dispatched_qty,
                    r.balance_qty,
                    r.pct_complete,
                    r.status,
                ]
                for r in rows
            ],
        )
    rows = db.query(ErpOrder).filter(ErpOrder.report_date == day).all()
    return (
        ["Order", "Customer", "Order date", "Required date", "Status", "Value"],
        [
            [r.order_no, r.customer, r.order_date, r.required_date, r.status, r.total_value]
            for r in rows
        ],
    )


def _xlsx(headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales report"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F35")
    for row in rows:
        ws.append(row)
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = max(
            14, len(header) + 3
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf(headers: list[str], rows: list[list], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24
    )
    styles = getSampleStyleSheet()
    data = [[Paragraph(escape(str(h)), styles["BodyText"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(cell)), styles["BodyText"]) for cell in row])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173f35")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d9e2dc")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8f5")]),
            ]
        )
    )
    doc.build([Paragraph(title, styles["Title"]), Spacer(1, 10), table])
    return out.getvalue()


@router.get("/export")
def export_report(
    dataset: str = Query(default="companies", pattern="^(companies|daily|dispatch|orders)$"),
    format: str = Query(default="xlsx", pattern="^(xlsx|pdf)$"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    headers, rows = _export_rows(db, dataset)
    title = f"Kalika sales {dataset} report - {_day(db) or 'no data'}"
    log_event(
        db,
        action="sales.export",
        actor_user_id=user.id,
        resource_type="sales_report",
        resource_id=dataset,
        status="success",
        detail={"format": format, "rows": len(rows)},
        commit=True,
    )
    if format == "pdf":
        content = _pdf(headers, rows, title)
        media = "application/pdf"
        filename = f"kalika-{dataset}-{_day(db) or 'empty'}.pdf"
    else:
        content = _xlsx(headers, rows)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"kalika-{dataset}-{_day(db) or 'empty'}.xlsx"
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
