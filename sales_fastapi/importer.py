"""Contact file parsing and normalisation.

Supports CSV, JSON, XLSX (openpyxl), XLS (xlrd) and VCF (vCard). Normalisation
maps the many header variants found in real GCS exports onto the Contact model.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from typing import Any

from .services import classify_contact

SUPPORTED_SUFFIXES = (".csv", ".json", ".xlsx", ".xls", ".vcf")


def _xlsx_rows(raw: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration:
        return []
    records: list[dict[str, Any]] = []
    for row in rows:
        if row is None or all(cell is None for cell in row):
            continue
        record = {
            headers[index]: row[index]
            for index in range(min(len(headers), len(row)))
            if headers[index]
        }
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def _xls_rows(raw: bytes) -> list[dict[str, Any]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=raw)
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows < 2:
        return []
    headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
    records: list[dict[str, Any]] = []
    for row_index in range(1, sheet.nrows):
        record = {
            headers[col]: sheet.cell_value(row_index, col)
            for col in range(sheet.ncols)
            if headers[col]
        }
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


_VCARD_BLOCK = re.compile(r"BEGIN:VCARD(.*?)END:VCARD", re.IGNORECASE | re.DOTALL)


def _vcf_records(raw: bytes) -> list[dict[str, Any]]:
    text = raw.decode("utf-8", errors="ignore")
    records: list[dict[str, Any]] = []
    for block in _VCARD_BLOCK.findall(text):
        def first(prop: str) -> str:
            match = re.search(rf"^{prop}[^:]*:(.+)$", block, re.IGNORECASE | re.MULTILINE)
            return match.group(1).strip() if match else ""

        emails = re.findall(r"^EMAIL[^:]*:(.+)$", block, re.IGNORECASE | re.MULTILINE)
        phones = re.findall(r"^TEL[^:]*:(.+)$", block, re.IGNORECASE | re.MULTILINE)
        company = first("ORG")
        name = first("FN") or first("N")
        if not any([name, company, emails, phones]):
            continue
        records.append(
            {
                "name": name,
                "company": company,
                "email": (emails[0].strip() if emails else ""),
                "phone": (phones[0].strip() if phones else ""),
                "title": first("TITLE"),
            }
        )
    return records


def records_from_bytes(name: str, raw: bytes) -> list[dict[str, Any]]:
    lower = (name or "").lower()
    if lower.endswith(".json"):
        try:
            parsed = json.loads(raw.decode("utf-8", errors="ignore"))
        except json.JSONDecodeError:
            return []
        if isinstance(parsed, dict):
            return parsed.get("contacts", [parsed])
        if isinstance(parsed, list):
            return parsed
        return []
    if lower.endswith(".csv"):
        text = raw.decode("utf-8", errors="ignore")
        return list(csv.DictReader(io.StringIO(text)))
    if lower.endswith(".xlsx"):
        try:
            return _xlsx_rows(raw)
        except Exception:
            return []
    if lower.endswith(".xls"):
        try:
            return _xls_rows(raw)
        except Exception:
            return []
    if lower.endswith(".vcf"):
        return _vcf_records(raw)
    return []


def _pick(record: dict[str, Any], *keys: str) -> str:
    for key in keys:
        for record_key in record:
            if record_key and str(record_key).strip().lower() == key:
                value = record[record_key]
                if value not in (None, ""):
                    return str(value).strip()
    return ""


def normalise_record(record: dict[str, Any]) -> dict[str, Any] | None:
    email_address = _pick(
        record,
        "email", "email (company / division)", "e-mail", "mail", "email id", "email address",
        "contact email", "e mail", "emailid",
    )
    company = _pick(
        record, "company", "organisation", "organization", "company name", "firm", "org",
    )
    name = _pick(
        record, "name", "contact", "contact name", "contact person", "full name", "person",
        "first name",
    )
    if not email_address and not company and not name:
        return None
    verdict = classify_contact(email_address, company or name)
    return {
        "company": company,
        "name": name,
        "email": email_address or f"unknown-{hashlib.md5((company or name).encode()).hexdigest()[:12]}@placeholder.local",
        "phone": _pick(record, "phone", "contact (phone)", "mobile", "phone number", "contact number", "tel"),
        "address": _pick(record, "address", "address (ahmednagar unit)", "location", "city"),
        "linkedin_company": _pick(record, "linkedin", "linkedin (company / key scm profile)", "linkedin url"),
        "purchase_contact_name": _pick(record, "purchase / scm contact (name, role)", "purchase contact", "scm contact"),
        "purchase_contact_email": _pick(record, "purchase contact details (email / phone)", "purchase email"),
        "purchase_contact_role": _pick(record, "title", "designation", "role", "position"),
        "domain": verdict["domain"],
        "intent": verdict["intent"],
        "context": company or name,
    }
