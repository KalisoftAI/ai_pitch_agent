"""Update Kalika_Dashboard_Summary Excel with the app contacts-sync results, push to GCS."""
import json, os
import google.auth
from google.cloud import storage
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary_v2.xlsx")
if not os.path.exists(XLSX):
    XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary.xlsx")
GCS_DEST = "sales-metadata/Kalika_Dashboard_Summary.xlsx"

report = json.load(open(os.path.join(BASE, "contacts_sync_report.json"), encoding="utf-8"))

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
SUB_FONT = Font(bold=True, size=11, color="2E74B5")
OK_FONT = Font(bold=True, color="1E7B34")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = load_workbook(XLSX)

# ---- Sheet 1: update Overview KPI note
ws = wb["1_Dashboard_Overview"]
for row in ws.iter_rows(min_col=1, max_col=3):
    for cell in row:
        if cell.value == "TOTAL CONTACT RECORDS (all lists)":
            ws.cell(row=cell.row, column=3, value=str(ws.cell(row=cell.row, column=3).value) +
                    " | SYNCED INTO APP DB (sales_pipeline.db): see sheet 2 sync block")
r = ws.max_row + 2
ws.cell(row=r, column=1, value="CONTACTS IN APP DB (after GCS sync)").font = Font(bold=True)
ws.cell(row=r, column=2, value=report["db_total_contacts"]).font = OK_FONT
ws.cell(row=r, column=3, value="user ai.solutions@kalisoftai.in | synced from sales-contacts/ via app importer pipeline (normalise_record), 30-Oct-2026")

# ---- Sheet 2: append per-sheet sync report
ws = wb["2_Contacts_Summary"]
r = ws.max_row + 2
ws.cell(row=r, column=1, value="APP CONTACTS SYNC — 30-Oct-2026 (into current app DB via /api/contacts/import/gcs pipeline; details preserved as-is: unmapped source columns kept verbatim in contact notes)").font = SUB_FONT
r += 1
headers = ["GCS source (file#sheet)", "Rows parsed", "Imported (new)", "Duplicates skipped", "No contact identity (skipped)", "Existing records field-filled"]
for j, h in enumerate(headers):
    c = ws.cell(row=r, column=1 + j, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = THIN
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
r += 1
for f in report["files"]:
    if "error" in f:
        vals = [f["gcs_path"], "ERROR: " + f["error"], "", "", "", ""]
    else:
        vals = [f["gcs_path"], f["rows_parsed"], f["imported"], f["skipped_duplicate"],
                f["skipped_no_contact_info"], f.get("existing_updated_empty_fields", 0)]
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=1 + j, value=v)
        c.border = THIN; c.alignment = WRAP
    r += 1
t = report["totals"]
for j, v in enumerate(["TOTAL", "", t["imported"], t["skipped_duplicate"], t["skipped_no_contact_info"], t.get("existing_updated_empty_fields", 0)]):
    c = ws.cell(row=r, column=1 + j, value=v)
    c.font = Font(bold=True); c.border = THIN
r += 2
notes = [
    f"App DB total contacts after sync: {report['db_total_contacts']} (was 7,924 -> +871 new from the 3 GCS contact files).",
    "Nothing added/removed: app standard field mapping used; extras (Call status, Requirements, Industry, Rates remarks, etc.) preserved verbatim in notes.",
    "Duplicates detected by app rule (user + email); orphan attribute-only rows (no name/company/email) are not contacts and were skipped.",
    "Field-fill: empty phone/address on 275 existing records filled from the same GCS rows (no invented data).",
]
for n in notes:
    ws.cell(row=r, column=1, value="• " + n).font = Font(italic=True, size=9)
    r += 1
for j, w in enumerate([52, 12, 13, 15, 17, 16], start=1):
    ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 10, w)

wb.save(XLSX)
print("Excel updated:", XLSX)

creds, _ = google.auth.default()
client = storage.Client(credentials=creds)
bucket = client.bucket("kalika_enterprises")
blob = bucket.blob(GCS_DEST)
blob.upload_from_filename(XLSX)
print(f"Pushed to gs://kalika_enterprises/{GCS_DEST} ({os.path.getsize(XLSX)} bytes)")
