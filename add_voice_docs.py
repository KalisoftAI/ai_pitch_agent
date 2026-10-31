"""Voice/SLM docs: create the dedicated markdown, append §20 to the main doc,
add Excel sheet 15, and push everything to GCS."""
import json
import os
import google.auth
from google.cloud import storage
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(BASE, "Kalika_Voice_Transcription_SLM.md")
MAIN = os.path.join(BASE, "gcs_data", "sales-metadata__Kalika_Real_Time_Customer_Feedback_System.md")
XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary_v2.xlsx")

VOICE_DOC = r"""# Kalika Voice Transcription & Gemma SLM Pipeline (30-Oct-2026)

> Backend job + API-only surface for call-recording intelligence on
> `gs://kalika_enterprises/call_recording/Vaishnavi-samp/`.
> Provider chain: **Gemma SLM (Ollama) → Gemini 2.5 Flash → deterministic heuristics**.
> Human review stays on for anything customer-facing. ERP remains untouched/prod.

---

## 1. Layered prompt pipeline (ASCII)

```text
 GCS: call_recording/Vaishnavi-samp/  (3 files: 60s m4a, 76s wav, 32s wav)
        │
        v
 L0  AUDIO PROBE      ffprobe + ffmpeg volumedetect/silencedetect
     duration · sample rate · channels · mean/max dB · silence ratio
        │
        v
 L1  TRANSCRIBE       Gemini 2.5 Flash (audio inline) → language + turns
     (needs GEMINI_API_KEY)                    │
        │                                      └─ no key → status pending_ai
        v
 L2  CONTEXT          products · quantities · rates · objections · summary
 L3  TONE             sentiment · urgency · politeness · interest · noise · pace
 L4  LEAD EXTRACT     name/company/phone · intent · products · next action · template
        │
        ├── provider chain: gemma (Ollama) → gemini → heuristic  (mode recorded)
        v
 voice_transcripts  ──> POST /api/voice/provision/{id}
        │                 contact + journey event (stage from intent)
        v
 DASHBOARD / JOURNEY   stories + sentiment/intent mix + recommended template
```

## 2. Provider chain & enablement

| Provider | Needs | Notes |
|---|---|---|
| **Gemma SLM (Ollama)** | `ollama serve` + `ollama pull gemma-4-27b-it` (or configured tag) | preferred (`VOICE_SLM_FIRST=true`), data stays on-device, zero cost |
| Gemini 2.5 Flash | `GEMINI_API_KEY` | required for L1 audio transcription; also fallback for L2–L4 |
| Heuristic | nothing | deterministic; always available; used for CI/tests |

Current runtime: `ollama reachable=false`, `gemini configured=false` → L0 runs for real,
rows stored as `pending_ai` with a clear error, and any ingested transcript is processed
heuristically (verified: lead + tone + provisioning all work).

## 3. API (API-only; no UI required)

| Method | Endpoint | Purpose |
|---|---|---|
| GET | `/api/voice/status` | prefix, file counts, SLM/provider status |
| GET | `/api/voice/files` | list audio + processed flag |
| GET | `/api/voice/providers` | chain order + reachability |
| POST | `/api/voice/transcribe` | run job (`{"file":"all","force":true}`) |
| POST | `/api/voice/ingest` | external ASR transcript → L2–L4 |
| GET | `/api/voice/transcripts` | list results (tone/lead) |
| GET | `/api/voice/transcripts/{id}` | full record incl. audio facts + layers |
| POST | `/api/voice/provision/{id}` | create/refresh contact + journey event |
| GET | `/api/voice/stats` | totals, audio minutes, sentiment/intent mix |

## 4. Scheduler

| Job | IST | UTC |
|---|---|---|
| `voice_transcription_sync` | 09:00 | `30 3 * * *` |

Executor: `python scripts/voice_sync.py [--force]` (same code path as the API).
18 cron tasks registered; the job uses the Gemma-first chain automatically.

## 5. Current numbers (live)

| File | Duration | Status |
|---|---|---|
| AUD-20260109-WA0005.m4a | 60.05 s | pending_ai (audio probed) |
| record-1793186132125.wav | 76.22 s | pending_ai (audio probed) |
| record-1793187022137.wav | 31.62 s | pending_ai (audio probed) |

Demo ingest → lead: phone 9822000000, products [VCI bags], 500 kg, intent `order`,
tone positive/high urgency → provisioning created contact **8796**, journey stage `evaluation`,
recommended template **WA-ORD-01**.

## 6. Enable Gemma SLM (deploy steps)

```bash
# on the host / Cloud Run sidecar
ollama serve &
ollama pull gemma-4-27b-it
# env
VOICE_SLM_FIRST=true
VOICE_SLM_MODEL=gemma-4-27b-it
LLM_LOCAL_BASE_URL=http://localhost:11434
# optional cloud: GEMINI_API_KEY=...  (needed for L1 transcription of audio)
```

Then `python scripts/voice_sync.py --force` or `POST /api/voice/transcribe`.
"""

STORE_SECTION = r"""

---

## 20. Voice Transcription & Gemma SLM Pipeline — 30-Oct-2026

Dedicated doc: `sales-metadata/Kalika_Voice_Transcription_SLM.md`.

- **Source:** `gs://kalika_enterprises/call_recording/Vaishnavi-samp/` (3 files, 167.9 s total).
- **Layers:** L0 audio probe (ffmpeg: duration/volume/silence) → L1 transcribe (Gemini flash audio)
  → L2 context → L3 tone → L4 lead extraction.
- **Provider chain:** **Gemma SLM (Ollama) → Gemini 2.5 Flash → heuristic**, mode recorded per record.
- **API-only:** `/api/voice/{status,files,providers,transcribe,ingest,transcripts,stats}` +
  `POST /api/voice/provision/{id}` (contact + journey event).
- **Scheduler:** `voice_transcription_sync` 09:00 IST (`30 3 * * *`); executor `scripts/voice_sync.py`.
- **Live:** audio probed for all 3 files (`pending_ai` until `GEMINI_API_KEY`), demo ingest produced
  lead (VCI bags, 500 kg, phone 9822000000), tone positive/high urgency, contact 8796 + journey
  stage `evaluation`, template `WA-ORD-01`.
- **Enable Gemma:** run Ollama + `ollama pull gemma-4-27b-it`, set `VOICE_SLM_FIRST=true`,
  `VOICE_SLM_MODEL=gemma-4-27b-it`, `LLM_LOCAL_BASE_URL`, and `GEMINI_API_KEY` for audio transcription.
- **Tests:** 14 voice tests (probe, layers, Gemma chain, fallbacks, ingest, provisioning, stats);
  full backend suite: **142 passing**.
"""

with open(DOC, "w", encoding="utf-8") as fh:
    fh.write(VOICE_DOC)
print("voice doc written:", DOC)

with open(MAIN, encoding="utf-8") as fh:
    content = fh.read()
marker = "\n---\n\n## 20. Voice Transcription"
if marker in content:
    content = content.split(marker)[0]
content = content.rstrip() + STORE_SECTION
with open(MAIN, "w", encoding="utf-8") as fh:
    fh.write(content)
print("main md updated:", len(content), "chars")

# ------------------------------------------------ Excel sheet 15
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(bold=True, size=11, color="2E74B5")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = load_workbook(XLSX)
if "15_Voice_SLM" in wb.sheetnames:
    del wb["15_Voice_SLM"]
ws = wb.create_sheet("15_Voice_SLM")


def table(r, headers, rows):
    for j, h in enumerate(headers):
        c = ws.cell(row=r, column=1 + j, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.border = THIN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = THIN; c.alignment = WRAP
        r += 1
    return r + 1


ws.cell(row=1, column=1, value="VOICE TRANSCRIPTION & GEMMA SLM PIPELINE — 30-Oct-2026").font = TITLE_FONT
ws.cell(row=2, column=1, value="gs://kalika_enterprises/call_recording/Vaishnavi-samp/ | chain: gemma (Ollama) -> gemini-2.5-flash -> heuristic | doc: sales-metadata/Kalika_Voice_Transcription_SLM.md").font = Font(italic=True, size=9, color="808080")
r = 4

r = table(r, ["Item", "Value", "Notes"], [
    ["Audio files", 3, "AUD-20260109-WA0005.m4a (60.05s) · record-1793186132125.wav (76.22s) · record-1793187022137.wav (31.62s)"],
    ["Audio minutes", 2.8, "sum 167.9 s"],
    ["Status", "pending_ai", "audio probed (L0) for all 3; L1 needs GEMINI_API_KEY"],
    ["Provider chain", "gemma -> gemini -> heuristic", "VOICE_SLM_FIRST=true, VOICE_SLM_MODEL=gemma-4-27b-it"],
    ["Gemma (Ollama) reachable", "false", "ollama serve + ollama pull gemma-4-27b-it"],
    ["Gemini configured", "false", "GEMINI_API_KEY required for audio transcription"],
    ["Scheduler", "voice_transcription_sync 30 3 * * *", "09:00 IST; executor scripts/voice_sync.py"],
    ["Tests", "14 voice / 142 total", "full backend suite passing"],
])

ws.cell(row=r, column=1, value="LAYERS").font = SUB_FONT; r += 1
r = table(r, ["Layer", "What it extracts", "Implementation"], [
    ["L0 audio probe", "duration, sample rate, channels, mean/max dB, silence ratio", "ffprobe + ffmpeg (always runs)"],
    ["L1 transcribe", "language, verbatim transcript, speaker turns", "Gemini 2.5 Flash audio inline"],
    ["L2 context", "products, quantities, rates, objections, commitments, summary", "Gemma SLM / Gemini / heuristic"],
    ["L3 tone", "sentiment, urgency, politeness, interest, noise, speech pace", "Gemma SLM / Gemini / heuristic"],
    ["L4 lead", "name/company/phone, intent, products, qty, next action, template, confidence", "Gemma SLM / Gemini / heuristic"],
])

ws.cell(row=r, column=1, value="API (API-ONLY)").font = SUB_FONT; r += 1
r = table(r, ["Method", "Endpoint", "Purpose"], [
    ["GET", "/api/voice/status", "prefix, file counts, SLM/provider status"],
    ["GET", "/api/voice/files", "list audio + processed flag"],
    ["GET", "/api/voice/providers", "chain order + reachability"],
    ["POST", "/api/voice/transcribe", "run job (file=all, force)"],
    ["POST", "/api/voice/ingest", "external ASR transcript -> L2-L4"],
    ["GET", "/api/voice/transcripts[/{id}]", "results incl. audio facts + layers"],
    ["POST", "/api/voice/provision/{id}", "contact + journey event creation"],
    ["GET", "/api/voice/stats", "totals, minutes, sentiment/intent mix"],
])

ws.cell(row=r, column=1, value="DEMO PROVISION RESULT").font = SUB_FONT; r += 1
r = table(r, ["Field", "Value"], [
    ["Contact created", "8796 (Shree Traders / phone 9822000000)"],
    ["Lead intent / products", "order / VCI bags, 500 kg"],
    ["Tone", "positive · urgency high · interest high"],
    ["Journey stage", "evaluation"],
    ["Recommended template", "WA-ORD-01"],
])
for i, w in enumerate([28, 40, 70], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

order = ["1_Dashboard_Overview", "2_Contacts_Summary", "3_Call_Status_Followups", "4_Sales_2026_Monthly",
         "5_Top_Customers_Products", "6_ERP_Dispatch_Orders", "7_Product_Rate_List", "8_WhatsApp_Data",
         "9_API_Details_Features", "10_Data_Sources_GCS", "11_Crons_Sync", "12_Graph_Preprocessing",
         "13_ERP_Sync_Forecast", "14_Sales_Deploy_UAT", "15_Voice_SLM"]
wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
wb.save(XLSX)
print("excel updated:", XLSX)

creds, _ = google.auth.default()
bucket = storage.Client(credentials=creds).bucket("kalika_enterprises")
bucket.blob("sales-metadata/Kalika_Voice_Transcription_SLM.md").upload_from_filename(DOC)
bucket.blob("sales-metadata/Kalika_Real_Time_Customer_Feedback_System.md").upload_from_filename(MAIN)
bucket.blob("sales-metadata/Kalika_Dashboard_Summary.xlsx").upload_from_filename(XLSX)
print("pushed voice doc + main md + excel")
