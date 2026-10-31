"""Add 'Crons & Sync' + 'Graph Preprocessing' sheets to the dashboard Excel; push to GCS."""
import json, os
import google.auth
from google.cloud import storage
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary_v2.xlsx")
GCS_DEST = "sales-metadata/Kalika_Dashboard_Summary.xlsx"

dump = json.load(open(os.path.join(BASE, "gcs_data", "interaction_dump.json"), encoding="utf-8"))
gstats = json.load(open(os.path.join(BASE, "gcs_data", "graph", "graph_stats.json"), encoding="utf-8"))

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(bold=True, size=11, color="2E74B5")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = load_workbook(XLSX)
for n in ["11_Crons_Sync", "12_Graph_Preprocessing"]:
    if n in wb.sheetnames:
        del wb[n]


def table(ws, r, headers, rows):
    for j, h in enumerate(headers):
        c = ws.cell(row=r, column=1 + j, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.border = THIN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = THIN; c.alignment = WRAP
        r += 1
    return r + 1


# ------------------------------------------------ 11. crons & sync
ws = wb.create_sheet("11_Crons_Sync")
ws.cell(row=1, column=1, value="BACKEND CRONS & DATA SYNC — UPDATED 30-Oct-2026").font = TITLE_FONT
ws.cell(row=2, column=1, value="POST /api/cron/run triggers a manual sweep of all registered tasks; scheduler backend: in-memory (Redis optional) | Cloud Scheduler hits /api/scheduler/run in production").font = Font(italic=True, size=9, color="808080")
r = table(ws, 4, ["Cron task", "Schedule", "Description", "Status"], [
    ["sync_gcs_contacts", "0 */6 * * *", "Import contacts from GCS bucket", "active"],
    ["refresh_email_connections", "*/30 * * * *", "Re-test IMAP/SMTP connections", "active"],
    ["linkedin_hiring_alerts", "0 8 * * *", "Daily LinkedIn hiring-keyword sweep", "active"],
    ["reddit_hiring_alerts", "0 8 * * *", "Daily Reddit hiring-thread sweep", "active"],
    ["youtube_signal_scan", "0 9 * * *", "YouTube company/tech signal scan", "future"],
    ["whatsapp_reminders", "0 10 * * *", "Queue follow-up WhatsApp reminders", "active"],
    ["events_whatsapp_due", "*/15 * * * *", "Dispatch queued event WhatsApp messages past schedule", "active"],
    ["pipeline_feedback_due", "0 * * * *", "NEW — send-due pass over feedback-pipeline queue (recent orders / new customers / leads)", "NEW 30-Oct"],
    ["build_interaction_dump", "0 2 * * *", "NEW — parse WhatsApp chat exports (kalika backup/) + app outbox -> interaction_dump.json with attachment details", "NEW 30-Oct"],
    ["rebuild_graph_preprocessing", "30 2 * * *", "NEW — networkx graph preprocessing -> nodes/edges/gexf (future graph DB)", "NEW 30-Oct"],
    ["kpi_snapshot", "0 23 * * *", "Nightly business-KPI snapshot", "active"],
])
s = dump["stats"]
ws.cell(row=r, column=1, value="INTERACTION DUMP — LAST RUN (build_interaction_dump)").font = SUB_FONT
r += 1
r = table(ws, r, ["Metric", "Value"], [
    ["Chats parsed (kalika backup/ zips)", s["chats"]],
    ["Total messages", s["messages"]],
    ["Unique participants", s["participants_total"]],
    ["Attachments catalogued (name + size + media type)", s["attachments"]],
    ["Media by type", " | ".join(f"{k}: {v}" for k, v in s["media_by_type"].items())],
    ["Media-omitted references in texts", s["media_omitted_refs"]],
    ["Outbound interactions (app whatsapp_messages)", s["mails"]],
    ["Output", "gcs_data/interaction_dump.json (messages + attachments + stats)"],
    ["Tests", "tests/test_graph_and_dump.py — 10 tests passing (chat parse, attachments, mail records, phone norm, templates, graph)"],
])
for i, w in enumerate([34, 16, 70, 14], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ------------------------------------------------ 12. graph preprocessing
ws = wb.create_sheet("12_Graph_Preprocessing")
ws.cell(row=1, column=1, value="GRAPH PRE-PROCESSING (networkx) — FUTURE GRAPH-Database SCOPE").font = TITLE_FONT
ws.cell(row=2, column=1, value="sales_fastapi/graph_build.py | exports: gcs_data/graph/{nodes.json, edges.json, graph.gexf, graph_stats.json} | GEXF imports into Gephi/Neo4j").font = Font(italic=True, size=9, color="808080")
r = table(ws, 4, ["Metric", "Value"], [
    ["Nodes total", gstats["nodes"]],
    ["Edges total", gstats["edges"]],
    ["Density", gstats["density"]],
    ["Connected components (weak)", gstats["connected_components"]],
    ["Generated at", gstats["generated_at"]],
])
ws.cell(row=r, column=1, value="NODES BY KIND").font = SUB_FONT; r += 1
r = table(ws, r, ["Node kind", "Count", "Meaning"], [
    ["contact", gstats["nodes_by_kind"]["contact"], "contacts table rows (8,795 after GCS sync)"],
    ["domain", gstats["nodes_by_kind"]["domain"], "email/company domains (gmail.com 307, cummins.com 21 ...)"],
    ["intent", gstats["nodes_by_kind"]["intent"], "general / sales 65 / procurement 68 / support 55 / hiring 28 / partnership"],
    ["chat_thread", gstats["nodes_by_kind"]["chat_thread"], "WhatsApp export chats from kalika backup/"],
    ["erp_customer", gstats["nodes_by_kind"]["erp_customer"], "top customers from 2026 sales registers"],
])
ws.cell(row=r, column=1, value="EDGES BY KIND").font = SUB_FONT; r += 1
r = table(ws, r, ["Edge kind", "Count", "Meaning"], [
    ["CONTACT_HAS_INTENT", gstats["edges_by_kind"]["CONTACT_HAS_INTENT"], "contact -> intent classification"],
    ["CONTACT_IN_DOMAIN", gstats["edges_by_kind"]["CONTACT_IN_DOMAIN"], "contact -> domain bucket"],
    ["ORDER_BY_CONTACT", gstats["edges_by_kind"]["ORDER_BY_CONTACT"], "ERP customer -> contact (company match)"],
    ["CHAT_MAPS_TO_CONTACT", gstats["edges_by_kind"].get("CHAT_MAPS_TO_CONTACT", 0), "chat thread -> contact (phone match)"],
])
ws.cell(row=r, column=1, value="TOP CONNECTED NODES").font = SUB_FONT; r += 1
r = table(ws, r, ["Node", "Degree"], [[t["node"], t["degree"]] for t in gstats["top_connected"]])
ws.cell(row=r, column=1, value="FUTURE SCOPE: load nodes/edges into a graph DB (Neo4j/Neptune) via graph.gexf; add MESSAGE_ABOUT + REPLIED_TO edges from reply threads; centrality-based lead scoring; community detection for account clusters.").font = Font(italic=True, size=9)
for i, w in enumerate([30, 14, 74], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# order sheets 1..12
order = ["1_Dashboard_Overview", "2_Contacts_Summary", "3_Call_Status_Followups", "4_Sales_2026_Monthly",
         "5_Top_Customers_Products", "6_ERP_Dispatch_Orders", "7_Product_Rate_List", "8_WhatsApp_Data",
         "9_API_Details_Features", "10_Data_Sources_GCS", "11_Crons_Sync", "12_Graph_Preprocessing"]
wb._sheets = [wb[n] for n in order if n in wb.sheetnames]

wb.save(XLSX)
print("Excel updated:", XLSX)

creds, _ = google.auth.default()
client = storage.Client(credentials=creds)
bucket = client.bucket("kalika_enterprises")
blob = bucket.blob(GCS_DEST)
blob.upload_from_filename(XLSX)
print(f"Pushed to gs://kalika_enterprises/{GCS_DEST} ({os.path.getsize(XLSX)} bytes)")
