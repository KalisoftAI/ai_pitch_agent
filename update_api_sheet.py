"""Replace sheet 9 with full CURRENT API catalog (FastAPI + Django) and push Excel to GCS."""
import os
import google.auth
from google.cloud import storage
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary.xlsx")
GCS_DEST = "sales-metadata/Kalika_Dashboard_Summary.xlsx"

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(bold=True, size=11, color="2E74B5")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = load_workbook(XLSX)
if "9_API_Details_Features" in wb.sheetnames:
    del wb["9_API_Details_Features"]
ws = wb.create_sheet("9_API_Details_Features")


def write_table(r, headers, rows):
    for j, h in enumerate(headers):
        c = ws.cell(row=r, column=1 + j, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.border = THIN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = THIN; c.alignment = WRAP
        r += 1
    return r + 1


r = 1
ws.cell(row=r, column=1, value="CURRENT API — DETAILS, FEATURES & DESCRIPTIONS").font = TITLE_FONT
r += 1
ws.cell(row=r, column=1, value="Live service: 'Kalisoft AI Sales Pipeline' v1.0.0 (FastAPI) — Cloud Run service 'kalisoft-sales', region asia-south1, all routes under /api | Legacy Django web app under /app/ | Updated 30-Oct-2026").font = Font(italic=True, size=9, color="808080")
r += 2

# ---------------------------------------------------------------- overview
ws.cell(row=r, column=1, value="A. SERVICE OVERVIEW & SECURITY").font = SUB_FONT
r += 1
r = write_table(r,
    ["Item", "Detail"],
    [
        ["Service name / version", "Kalisoft AI Sales Pipeline — v1.0.0 (FastAPI)"],
        ["Deployment", "Google Cloud Run: service 'kalisoft-sales', region asia-south1 (cloudbuild.fastapi.yaml); Artifact Registry image; separate sales & checks pipelines"],
        ["API base path", "/api (14 routers); interactive docs at /docs & /redoc in non-production"],
        ["Auth", "Google OAuth login -> signed token (POST /api/auth/google); per-user isolation via get_current_user dependency"],
        ["Middleware / guardrails", "TrustedHostMiddleware (host allow-list), RequestGuardrailsMiddleware (request safety/PII guardrails), CORS restricted to configured origins, methods GET/POST/PATCH/DELETE"],
        ["Data stores", "SQL DB via Alembic migrations (contacts, outreach, feedback, google mailboxes, events); Google Sheets sync; GCS import; Gmail/IMAP/SMTP; WhatsApp gateway (wechaty)"],
        ["Governance", "Audit events, retention purge, usage metering, GCP policy checks, LLM task plans & guardrailed runs"],
    ])

# ---------------------------------------------------------------- FastAPI catalog
ws.cell(row=r, column=1, value="B. FASTAPI ENDPOINT CATALOG (current, implemented)").font = SUB_FONT
r += 1
r = write_table(r,
    ["Module (router)", "Method", "Endpoint", "Description", "Key Features", "Application / Use"],
    [
        # system
        ["system", "GET", "/api/config", "Public runtime config", "Exposes non-secret client config", "Frontend bootstrap"],
        ["system", "GET", "/api/health", "Health check", "DB-ready flag from lifespan init", "Uptime monitoring / Cloud Run probes"],
        ["system", "GET", "/api/integrations", "List integration status", "Shows connected external services", "Integrations dashboard"],
        ["system", "POST", "/api/cron/run", "Trigger cron tasks manually", "Auth-protected scheduler kick-off", "Nightly jobs, catch-up runs"],
        # auth
        ["auth", "POST", "/api/auth/google", "Login with Google", "Verifies Google ID token, issues session token", "User sign-in"],
        ["auth", "GET", "/api/auth/me", "Current user profile", "Returns authenticated user", "Session check, UI identity"],
        # contacts
        ["contacts", "GET", "/api/contacts", "List contacts", "Filters: domain, intent; per-user", "Contact lists, segmentation"],
        ["contacts", "POST", "/api/contacts", "Create contact", "Validated create (201)", "Manual lead entry"],
        ["contacts", "GET", "/api/contacts/stats", "Contact statistics", "Counts by domain/intent/status", "Dashboard KPI cards"],
        ["contacts", "GET/PATCH/DELETE", "/api/contacts/{id}", "Get / update / delete one contact", "Full CRUD on single record", "Contact detail editing"],
        ["contacts", "POST", "/api/contacts/upload", "Bulk upload contacts", "Multi-file upload (CSV/XLSX)", "Importing marketing lists (e.g., Pune/marketing lists)"],
        ["contacts", "POST", "/api/contacts/import/gcs", "Import contacts from GCS", "Optional bucket prefix filter", "Direct sync from kalika_enterprises bucket"],
        # email
        ["email (IMAP/SMTP)", "GET/POST", "/api/email/connections", "List / add mail connections", "Per-user SMTP/IMAP creds", "Mailbox management"],
        ["email (IMAP/SMTP)", "POST", "/api/email/connections/{id}/test", "Test a connection", "Validates login + deliverability", "Setup verification"],
        ["email (IMAP/SMTP)", "POST", "/api/email/connections/{id}/fetch", "Fetch mail via IMAP", "Filterable fetch (EmailFetchIn)", "Lead reply ingestion"],
        ["email (IMAP/SMTP)", "GET", "/api/email/connections/{id}/inbox", "Read inbox", "limit param (default 15)", "Inbox preview in app"],
        ["email (IMAP/SMTP)", "POST", "/api/email/send", "Send email", "Send via chosen connection_id", "Outreach sending"],
        # google email
        ["email/google (Gmail OAuth)", "GET", "/api/email/google/mailboxes", "List linked Gmail mailboxes", "OAuth-connected accounts", "Multi-mailbox sending"],
        ["email/google (Gmail OAuth)", "GET", "/api/email/google/authorize", "Start Gmail OAuth", "Returns consent URL", "Connect new Gmail"],
        ["email/google (Gmail OAuth)", "GET", "/api/email/google/callback", "OAuth callback", "Stores tokens (encrypted)", "Completes Gmail link"],
        ["email/google (Gmail OAuth)", "POST", "/api/email/google/mailboxes/{id}/fetch", "Fetch Gmail messages", "Gmail API fetch with filters", "Reply/thread sync"],
        # social
        ["social", "GET/POST", "/api/social/linkedin/profiles | /search", "LinkedIn profiles & search", "Stores found profiles", "Prospect discovery"],
        ["social", "GET/POST", "/api/social/reddit/posts | /search", "Reddit posts & search", "Hiring/intent alerts", "Intent signals"],
        ["social", "POST", "/api/social/youtube/search", "YouTube search", "Channel/video discovery", "Research & prospecting"],
        ["social", "POST/GET", "/api/social/whatsapp/messages", "Queue + list WhatsApp messages", "Outbound queue (201 on create)", "WhatsApp follow-ups"],
        # whatsapp
        ["whatsapp", "GET", "/api/whatsapp/status", "Gateway status", "wechaty-gateway health", "Channel monitoring"],
        ["whatsapp", "GET", "/api/whatsapp/messages", "List WhatsApp messages", "limit param (default 50)", "Message history"],
        ["whatsapp", "POST", "/api/whatsapp/send", "Send WhatsApp message", "Guardrailed send via gateway", "Customer follow-up messaging"],
        ["whatsapp", "POST", "/api/whatsapp/webhook", "Inbound WhatsApp webhook", "Receives incoming messages/events", "Two-way conversations, replies capture"],
        # templates
        ["templates", "GET", "/api/templates/meta", "Template metadata", "Channels/types available", "Template UI options"],
        ["templates", "GET/POST", "/api/templates", "List / create templates", "Channel filter; per-user", "Message template library"],
        ["templates", "PATCH/DELETE", "/api/templates/{id}", "Update / delete template", "Full CRUD", "Template maintenance"],
        ["templates", "POST", "/api/templates/seed", "Seed default templates", "Idempotent seeding", "First-run setup"],
        # campaigns
        ["campaigns", "POST", "/api/campaigns", "Create campaign (201)", "Generates per-recipient messages", "AI outreach campaigns"],
        ["campaigns", "GET", "/api/campaigns | /{id}", "List / get campaign", "CampaignDetail with messages", "Campaign management"],
        ["campaigns", "PATCH", "/api/campaigns/{id}/messages/{mid}", "Edit a campaign message", "Human-in-the-loop editing", "Review before send"],
        ["campaigns", "POST", "/api/campaigns/{id}/approve", "Approve campaign", "Approval gate before sending", "Compliance/quality control"],
        ["campaigns", "POST", "/api/campaigns/{id}/send", "Send approved campaign", "Returns send result stats", "Bulk outreach execution"],
        ["campaigns", "DELETE", "/api/campaigns/{id}", "Delete campaign (204)", "—", "Cleanup"],
        # scheduler + kpi
        ["scheduler", "POST/GET", "/api/scheduler/run | /status | /tasks | /clear", "Control & inspect scheduler", "Run now, view tasks, clear queue", "Automation ops"],
        ["kpi", "GET", "/api/kpi/overview", "KPI overview", "Headline funnel metrics", "Dashboard top cards"],
        ["kpi", "GET", "/api/kpi/pipeline", "Pipeline KPIs", "Stage-wise counts", "Pipeline view"],
        ["kpi", "GET", "/api/kpi/outreach", "Outreach KPIs", "Sends/opens/replies", "Outreach performance"],
        ["kpi", "GET", "/api/kpi/costs", "Cost KPIs", "LLM/API spend tracking", "Cost control"],
        ["kpi", "GET", "/api/kpi/website", "Website KPIs", "Site analytics summary", "Marketing insights"],
        # feedback
        ["feedback", "POST", "/api/feedback", "Submit feedback (201)", "Guardrailed intake", "In-app feedback"],
        ["feedback", "GET", "/api/feedback", "List feedback", "Per-user listing", "Feedback review"],
        ["security", "GET", "/api/security/status", "Security status", "Guardrails/PII status", "Compliance check"],
        # governance
        ["governance", "GET", "/api/governance/policy", "Policy view", "Active governance rules", "Admin review"],
        ["governance", "GET", "/api/governance/gcp", "GCP posture", "GCP config checks", "Cloud compliance"],
        ["governance", "GET", "/api/governance/audit", "Audit events list", "AuditEventOut records", "Audit trail"],
        ["governance", "POST", "/api/governance/retention/purge", "Purge per retention policy", "Data lifecycle enforcement", "Privacy compliance"],
        ["governance", "GET", "/api/governance/usage", "Usage metering", "Per-user usage stats", "Quota/billing insight"],
        # llm
        ["llm", "GET", "/api/llm/tasks", "List LLM tasks", "Allowed task catalog", "AI feature discovery"],
        ["llm", "GET", "/api/llm/plan/{task}", "Get task plan", "Prompt/graph plan for a task", "Transparency of AI runs"],
        ["llm", "POST", "/api/llm/run", "Run an LLM task", "Guardrailed execution (LLMRunOut)", "AI content/analysis generation"],
        # events
        ["events", "GET", "/api/events/summary", "Event summary", "Aggregate counts", "Event dashboard"],
        ["events", "GET", "/api/events | /{id}", "List / event detail", "Search query 'q'", "Event management (e.g., Packvision expo)"],
        ["events", "GET", "/api/events/recipients", "List event recipients", "Searchable", "Invitee lists"],
        ["events", "POST", "/api/events/import", "Import recipients workbook", "Upload XLSX of attendees", "Expo lead import"],
        ["events", "GET", "/api/events/messages/all", "List event messages", "Filter by event_id", "Event comms tracking"],
        ["events", "POST", "/api/events/messages/{id}/send | /cancel", "Send or cancel one event message", "Per-message control", "Scheduled follow-ups"],
        ["events", "POST", "/api/events/messages/send-due", "Send all due messages", "limit param (default 50)", "Batch scheduled sending"],
        ["events", "GET", "/api/events/message-template-default", "Default event template", "—", "Template fallback"],
        ["events", "GET", "/api/events/{id}/message-template", "Per-event template", "Optional recipient preview", "Personalized invites"],
        ["events", "POST", "/api/events/{id}/schedule", "Schedule event messages", "EventScheduleIn payload", "Follow-up automation"],
    ])

# ---------------------------------------------------------------- Django legacy
ws.cell(row=r, column=1, value="C. DJANGO WEB-APP ENDPOINTS (legacy UI + AJAX, under /app/)").font = SUB_FONT
r += 1
r = write_table(r,
    ["App", "Endpoint(s)", "Description", "Key Features", "Application / Use"],
    [
        ["dashboard", "/app/api/contacts/list/ · /count/ · /add/ · /upload/ · /delete/ · /api/diagnostics/", "Contact list APIs for web dashboard", "List/count/add/upload/delete contacts; diagnostics endpoint", "Web contact management"],
        ["dashboard", "/app/api/csv/list/ · /count/ · /download/<category>/<file>/ · /delete/", "CSV file manager APIs", "GCS-backed CSV listing/download/delete", "Shared data files"],
        ["dashboard", "/app/gmail-settings/ · /api/gmail/save|test|status/", "Per-user Gmail SMTP settings", "Save creds, test send, status", "User mailbox setup"],
        ["scraper", "/app/scraper/search/profile/ · /search/domain/ · /analyze/ · /pitch/", "Profile/domain scraping & AI pitch", "Profile search, domain search, analyze, generate pitch", "Lead research"],
        ["scraper", "/app/scraper/save/ · /update/ · /delete/ · /list/ · /download-csv/", "Scraped-contacts CRUD", "Save to contacts, update, delete, list, CSV export", "Lead capture"],
        ["scraper", "/app/scraper/keywords/* · /locations/* · /runs/* · /run-now/ · /scheduler/run/ · /progress/ · /stats/", "Scheduled daily scraping", "Keyword & location watchlists, run history, scheduler, live progress, stats", "Automated lead harvesting"],
        ["extractor", "/app/search/ · /download_csv/ · /sync-all/", "Search extractor & Google Sheet sync", "Search, CSV download, full sync to Google Sheet", "Data extraction workflows"],
        ["ai_agent_pitch", "/app/pitch/save-template/ · load-template/<id>/ · delete-template/<id>/ · generate-subject/ · enhance-with-ai/", "Email pitch studio (AJAX)", "Template CRUD, AI subject generation, AI body enhancement (gemma_service)", "Email pitch creation"],
        ["ai_agent_pitch", "/app/pitch/get-campaigns/ · mark-opened/<cid>/<email>/ · dashboard/<cid>/", "Campaign tracking", "List campaigns, open-tracking pixel, campaign detail", "Engagement tracking"],
        ["ai_agent_pitch", "/app/pitch/campaign-engine/config(+/save/) · run/ · progress/ · runs/(+/detail/)", "AI Auto Campaign Engine", "Config save, run, live progress, run history/detail (auto_campaign_service)", "Automated campaigns"],
        ["pitch_generator", "/app/generator/create/ · /result/<id>/", "Pitch generator", "Create + view generated pitch", "Quick pitch docs"],
        ["auth (django)", "/login/ · /logout/ · /signup/ · /password-reset/…", "Web auth flows", "Django auth incl. password reset", "User accounts"],
    ])

# ---------------------------------------------------------------- planned
ws.cell(row=r, column=1, value="D. PLANNED — QR FEEDBACK SYSTEM APIs (spec: Kalika_Real_Time_Customer_Feedback_System.md)").font = SUB_FONT
r += 1
r = write_table(r,
    ["Planned API", "Purpose / Description", "Key Features", "Application"],
    [
        ["QR Code Service", "Unique QR entry points (general / delivery / product)", "Secure order refs; source identifiers; no personal data in URLs", "Invoices, challans, labels, brochures"],
        ["Customer Feedback API", "Collect ratings + comments in real time", "Conditional follow-up questions; delivery vs product feedback; duplicate protection", "Post-delivery & product feedback"],
        ["Complaint Ticket API", "Create/route/track tickets", "Open→Assigned→In Progress→Resolved→Closed/Reopened; first-response & resolution timestamps; photo attachments", "Issue resolution workflow"],
        ["ERP Lookup API", "Order/customer/product/delivery lookup from Kalika ERP", "Order/invoice references; live price/stock/delivery estimates", "Auto-identify order on scan"],
        ["WhatsApp Messaging (templates)", "7 ready flows: welcome, delivery, product, issue, catalogue, suggestion, thank-you", "Interactive buttons/lists; consent & opt-out; 1 reminder after 24–48h", "Customer engagement automation"],
        ["Catalogue API", "13 product categories with specs", "Images + enquiry + quotation request", "Mobile catalogue via QR"],
        ["Alerts API", "Low rating→quality; delivery complaint→dispatch; repeats→investigation; unassigned→supervisor", "Configurable alert rules", "Ops escalation"],
        ["Feedback Analytics API", "Dashboard KPIs & exports", "Avg rating, distribution, complaints by product/category, response & resolution times, QR scan vs submit rate", "Management dashboard"],
    ])

for col in range(1, 7):
    ws.column_dimensions[get_column_letter(col)].width = [24, 16, 44, 40, 44, 34][col - 1]

# save (fallback if the file is open in Excel)
try:
    wb.save(XLSX)
    print("Excel updated:", XLSX)
except PermissionError:
    XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary_v2.xlsx")
    wb.save(XLSX)
    print("Original file locked (open in Excel?). Saved as:", XLSX)

# ---------------------------------------------------------------- push to GCS
creds, _ = google.auth.default()
client = storage.Client(credentials=creds)
bucket = client.bucket("kalika_enterprises")
blob = bucket.blob(GCS_DEST)
blob.upload_from_filename(XLSX)
print(f"Pushed to gs://kalika_enterprises/{GCS_DEST} ({os.path.getsize(XLSX)} bytes)")
