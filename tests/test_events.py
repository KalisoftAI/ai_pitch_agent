"""Events dashboard: workbook import, showcase filters and WhatsApp scheduling."""

import io
from datetime import datetime, timedelta, timezone

from sales_fastapi.config import settings
from sales_fastapi.event_importer import normalise_phone, parse_event_when, parse_workbook
from sales_fastapi.routers.events import DEFAULT_EVENT_TEMPLATE


def test_parse_when_accepts_24_hour_times_without_meridiem():
    # Tracker rows such as "10:00 - 18:00" leave the am/pm group empty.
    parsed = parse_event_when("24 sept 2026, 10:00 - 18:00")
    assert parsed["starts_at"] == datetime(2026, 9, 24, 10, 0)
    assert parsed["ends_at"] == datetime(2026, 9, 24, 18, 0)


def test_parse_when_accepts_yearless_day_month_cells():
    # "15 dec" only matches the two-group day+month pattern.
    parsed = parse_event_when("15 dec")
    assert (parsed["starts_at"].month, parsed["starts_at"].day) == (12, 15)
    assert parsed["starts_at"].year in {datetime.utcnow().year, datetime.utcnow().year + 1}
    assert parse_event_when("12 Dec, 10:00")["starts_at"].hour == 10


def _workbook_bytes() -> bytes:
    """A miniature stand-in for ``events list.xlsx`` covering every layout.

    Dates are relative to today so the "upcoming" fixtures never go stale.
    """
    from openpyxl import Workbook

    today = datetime.utcnow().date()

    def day(offset: int) -> str:
        return f"{(today + timedelta(days=offset)).strftime('%d %B, %Y')}"

    workbook = Workbook()
    exhibitions = workbook.active
    exhibitions.title = "Sheet1"
    exhibitions.append(["Events ", "Date and Location", "About ", "To do list "])
    exhibitions.append(
        [
            "https://example.com/pi ecc",
            f"{day(30)} - {day(32)}\n\nPIECC, Moshi, Pune, India",
            "AI Digital Manufacturing Expo",
            "Send brochures",
        ]
    )
    exhibitions.append(
        ["https://example.com/webinar", f"{day(10)} 23:30 Online", "Agentic AI Webinar", ""]
    )

    contacts = workbook.create_sheet("Sheet6")
    contacts.append(["Source,Company,Name,Phone,Email_Notes"])
    contacts.append(["Note,Suncrest International,Vishal Deshmukh,9960099009,Material & HSN"])
    contacts.append(["Note,Akona Engineering,B Prasad,9561428584, 9970261210"])

    links = workbook.create_sheet("Sheet7")
    links.append(["Event 2026"])
    links.append(
        [
            "\nhttps://luma.com/abc",
            "Next Level AI Workshop\nEvent by CoLab Lisbon\n"
            f"{(today + timedelta(days=5)).strftime('%a, %b %d, %Y')}, 4:30 PM - 9:30 PM (your local time)\n"
            "Avenida Infante Dom Henrique 143, Beato, Lisboa",
        ]
    )

    tracker = workbook.create_sheet("Sheet8")
    tracker.append(["Some stray title"])
    tracker.append([None, None, None, None, None, None, None])
    tracker.append(
        [
            "Event name",
            "Arranged by",
            "Time",
            "Meeting link",
            "Attendance",
            "Meeting mode",
            "About",
        ]
    )
    tracker.append(
        [
            "AI Customer Support Summit",
            "Event by Customer Success Collective",
            f"{(today + timedelta(days=1)).strftime('%b %d, %Y')}, 12:30 PM - 2:30 PM",
            "\nhttps://example.com/summit",
            "19 attendees",
            "Online",
            "A practical summit",
        ]
    )
    workbook.create_sheet("Sheet9")  # empty placeholder sheet

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _import(client, headers, **data):
    return client.post(
        "/api/events/import",
        headers=headers,
        files={
            "file": (
                "events list.xlsx",
                _workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data=data,
    )


def _enable_wechaty(monkeypatch):
    monkeypatch.setattr(settings, "WHATSAPP_ENABLED", True)
    monkeypatch.setattr(settings, "WECHATY_GATEWAY_TOKEN", "gw-token")


# --------------------------------------------------------------------------- #
# parser
# --------------------------------------------------------------------------- #
def test_parse_when_handles_real_world_formats():
    assert parse_event_when("9 April, 2026 - 11 April, 2026\n\nPIECC, Pune")[
        "starts_at"
    ] == datetime(2026, 4, 9)
    assert parse_event_when("9 April, 2026 - 11 April, 2026")["ends_at"] == datetime(2026, 4, 11)
    assert parse_event_when("19-20-21 Feb, 2026,")["starts_at"] == datetime(2026, 2, 19)
    assert parse_event_when("15/01/2026 23:30:00 Online")["starts_at"] == datetime(
        2026, 1, 15, 23, 30
    )
    assert parse_event_when("5, 6, 7, 8 FEBRUARY - 2026")["starts_at"] == datetime(2026, 2, 1)
    assert parse_event_when("sept 24 -25 Denver colorado")["starts_at"] == datetime(2026, 9, 24)
    assert parse_event_when(datetime(2026, 9, 23, 15, 30))["starts_at"] == datetime(
        2026, 9, 23, 15, 30
    )
    assert parse_event_when("")["starts_at"] is None


def test_parse_when_extracts_location_below_the_date():
    parsed = parse_event_when("28 -30 January 2026\n\nPIECC, Moshi, Pune, India")
    assert parsed["starts_at"] == datetime(2026, 1, 28)
    assert "Pune" in parsed["location"]


def test_normalise_phone_splits_multi_number_cells():
    assert normalise_phone("+91 9769326919") == ["919769326919"]
    assert normalise_phone("9561428584, 9970261210") == ["919561428584", "919970261210"]
    assert normalise_phone("n/a") == []


def test_parse_workbook_handles_every_layout():
    parsed = parse_workbook(_workbook_bytes())
    titles = [event["title"] for event in parsed["events"]]
    assert "AI Digital Manufacturing Expo" in titles
    assert "Next Level AI Workshop" in titles
    assert "AI Customer Support Summit" in titles
    # Section headings and empty sheets are skipped.
    assert "Event 2026" not in titles
    assert len(parsed["events"]) == 4

    phones = {contact["phone"] for contact in parsed["contacts"]}
    assert phones == {"919960099009", "919561428584", "919970261210"}
    suncrest = next(c for c in parsed["contacts"] if c["name"] == "Vishal Deshmukh")
    assert suncrest["company"] == "Suncrest International"
    assert suncrest["notes"] == "Material & HSN"


# --------------------------------------------------------------------------- #
# import
# --------------------------------------------------------------------------- #
def test_import_handles_bare_24_hour_times(client, auth_headers):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Event link", "Date", "About", "To do list"])
    sheet.append(
        [
            "https://example.com/ai-day",
            "5 March 2027, 10:00 - 18:00",
            "Full day AI track",
            "",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = client.post(
        "/api/events/import",
        headers=auth_headers,
        files={"file": ("events list.xlsx", buffer.getvalue(), "application/vnd.ms-excel")},
    )
    assert response.status_code == 200, response.text
    events = client.get("/api/events", headers=auth_headers).json()
    assert len(events) == 1
    assert events[0]["starts_at"] == "2027-03-05T10:00:00"
    assert events[0]["ends_at"] == "2027-03-05T18:00:00"


def test_import_requires_auth(client):
    response = client.post("/api/events/import")
    assert response.status_code == 401


def test_import_uploads_events_and_contacts(client, auth_headers):
    response = _import(client, auth_headers)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["events_imported"] == 4
    assert body["contacts_imported"] == 3
    assert "Sheet1" in body["sheets"]

    events = client.get("/api/events", headers=auth_headers).json()
    assert len(events) == 4
    assert {event["mode"] for event in events} >= {"in_person", "online"}


def test_import_is_idempotent(client, auth_headers):
    first = _import(client, auth_headers).json()
    second = _import(client, auth_headers).json()
    assert second["events_imported"] == 0
    assert second["events_updated"] == first["events_imported"] + first["events_updated"]
    assert second["contacts_imported"] == 0
    assert len(client.get("/api/events", headers=auth_headers).json()) == 4


def test_import_can_skip_contacts(client, auth_headers):
    response = _import(client, auth_headers, include_contacts="false")
    assert response.status_code == 200
    assert response.json()["contacts_imported"] == 0
    assert client.get("/api/events/recipients", headers=auth_headers).json() == []


def test_import_rejects_non_workbook_payload(client, auth_headers):
    response = client.post(
        "/api/events/import",
        headers=auth_headers,
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )
    assert response.status_code == 400


def test_import_from_local_path_is_scoped_to_data_dir(client, auth_headers):
    response = client.post(
        "/api/events/import",
        headers=auth_headers,
        data={"source_path": "../../secrets.xlsx"},
    )
    assert response.status_code == 400


def test_import_missing_local_workbook_returns_404(client, auth_headers):
    response = client.post(
        "/api/events/import",
        headers=auth_headers,
        data={"source_path": "does-not-exist.xlsx"},
    )
    assert response.status_code == 404


# --------------------------------------------------------------------------- #
# showcase
# --------------------------------------------------------------------------- #
def test_events_list_requires_auth(client):
    assert client.get("/api/events").status_code == 401


def test_events_can_be_filtered(client, auth_headers):
    _import(client, auth_headers)
    online = client.get("/api/events?mode=online", headers=auth_headers).json()
    assert online and all(event["mode"] == "online" for event in online)

    searched = client.get("/api/events?q=manufacturing", headers=auth_headers).json()
    assert len(searched) == 1
    assert "Manufacturing" in searched[0]["title"]


def test_summary_counts_events_and_recipients(client, auth_headers):
    _import(client, auth_headers)
    body = client.get("/api/events/summary", headers=auth_headers).json()
    assert body["events"] == 4
    assert body["recipients"] == 3
    assert body["upcoming"] + body["past"] + body["unscheduled"] == 4
    assert body["last_imported_at"]


def test_events_are_tenant_scoped(client, auth_headers):
    _import(client, auth_headers)
    event_id = client.get("/api/events", headers=auth_headers).json()[0]["id"]

    other_token = client.post(
        "/api/auth/google", json={"id_token": "dev:other@kalisoftai.com"}
    ).json()
    other = {"Authorization": f"Bearer {other_token['access_token']}"}

    assert client.get(f"/api/events/{event_id}", headers=other).status_code == 404
    assert client.get("/api/events", headers=other).json() == []
    assert client.get("/api/events/recipients", headers=other).json() == []
    assert client.get("/api/events/summary", headers=other).json()["events"] == 0


# --------------------------------------------------------------------------- #
# templates and scheduling
# --------------------------------------------------------------------------- #
def _first_event(client, headers):
    return client.get("/api/events?status=upcoming", headers=headers).json()[0]


def test_message_template_renders_personalised_preview(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    recipient = client.get("/api/events/recipients", headers=auth_headers).json()[0]

    body = client.get(
        f"/api/events/{event['id']}/message-template?recipient_id={recipient['id']}",
        headers=auth_headers,
    ).json()
    assert recipient["name"] in body["preview"]
    assert "{{" not in body["preview"]
    assert body["recipient_count"] == 3


def test_default_template_never_leaks_optional_placeholders(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    # No recipient selected: name/company fall back instead of shipping "{{...}}".
    body = client.get(f"/api/events/{event['id']}/message-template", headers=auth_headers).json()
    assert body["preview_name"] == "there"
    assert "your organisation" in body["preview"]
    assert "{{" not in body["preview"]


def test_unknown_template_variable_is_rejected(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=auth_headers).json()]
    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={"recipient_ids": ids, "template": "Hi {{nope}}"},
    )
    assert response.status_code == 422
    assert "nope" in response.json()["detail"]


def test_schedule_requires_a_recipient(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={"recipient_ids": [], "template": "Hello"},
    )
    assert response.status_code == 422


def test_schedule_queues_messages_without_sending(client, auth_headers, monkeypatch):
    monkeypatch.setattr(settings, "WHATSAPP_ENABLED", False)
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=auth_headers).json()]

    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={"recipient_ids": ids, "template": DEFAULT_EVENT_TEMPLATE},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["queued"] == 3
    assert body["scheduled_at"]

    queued = client.get("/api/events/messages/all", headers=auth_headers).json()
    assert queued["counts"] == {"queued": 3}
    assert all(row["status"] == "queued" and row["sent_at"] is None for row in queued["items"])
    assert all("{{" not in row["body"] for row in queued["items"])


def test_schedule_is_idempotent_per_recipient(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=auth_headers).json()]
    payload = {"recipient_ids": ids, "template": DEFAULT_EVENT_TEMPLATE}

    first = client.post(f"/api/events/{event['id']}/schedule", headers=auth_headers, json=payload)
    second = client.post(f"/api/events/{event['id']}/schedule", headers=auth_headers, json=payload)
    assert first.json()["queued"] == 3
    assert second.json()["queued"] == 0
    assert second.json()["skipped"] == 3
    assert (
        client.get("/api/events/messages/all", headers=auth_headers).json()["counts"]["queued"] == 3
    )


def test_preview_only_does_not_queue(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={"recipient_ids": [], "template": DEFAULT_EVENT_TEMPLATE, "preview_only": True},
    )
    assert response.status_code == 200
    assert response.json()["queued"] == 0
    assert response.json()["preview"]
    assert client.get("/api/events/messages/all", headers=auth_headers).json()["items"] == []


def test_schedule_without_a_template_uses_the_default(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=auth_headers).json()]

    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={"recipient_ids": ids},
    )
    assert response.status_code == 200, response.text
    assert response.json()["queued"] == 3
    queued = client.get("/api/events/messages/all", headers=auth_headers).json()["items"]
    assert all("{{" not in row["body"] for row in queued)


def test_schedule_rejects_an_explicitly_empty_template(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=auth_headers).json()]
    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={"recipient_ids": ids, "template": "   "},
    )
    assert response.status_code == 422
    assert "empty" in response.json()["detail"]


def test_aware_schedule_is_stored_as_naive_utc(client, auth_headers):
    _import(client, auth_headers)
    event = _first_event(client, auth_headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=auth_headers).json()]

    ist = timezone(timedelta(hours=5, minutes=30))
    response = client.post(
        f"/api/events/{event['id']}/schedule",
        headers=auth_headers,
        json={
            "recipient_ids": ids,
            "template": DEFAULT_EVENT_TEMPLATE,
            "scheduled_at": datetime(2026, 10, 1, 12, 0, tzinfo=ist).isoformat(),
        },
    )
    assert response.status_code == 200, response.text
    # 12:00 IST is 06:30 UTC; the offset is applied exactly once.
    assert response.json()["scheduled_at"] == "2026-10-01T06:30:00"
    queued = client.get("/api/events/messages/all", headers=auth_headers).json()["items"]
    assert all(row["scheduled_at"] == "2026-10-01T06:30:00" for row in queued)


# --------------------------------------------------------------------------- #
# explicit send
# --------------------------------------------------------------------------- #
def _queue(client, headers, *, when=None):
    _import(client, headers)
    event = _first_event(client, headers)
    ids = [row["id"] for row in client.get("/api/events/recipients", headers=headers).json()]
    payload = {"recipient_ids": ids, "template": DEFAULT_EVENT_TEMPLATE}
    if when is not None:
        payload["scheduled_at"] = when.isoformat()
    client.post(f"/api/events/{event['id']}/schedule", headers=headers, json=payload)
    return event


def test_send_due_returns_503_when_whatsapp_disabled(client, auth_headers, monkeypatch):
    monkeypatch.setattr(settings, "WHATSAPP_ENABLED", False)
    _queue(client, auth_headers, when=datetime.utcnow() - timedelta(minutes=1))
    assert client.post("/api/events/messages/send-due", headers=auth_headers).status_code == 503


def test_send_due_dispatches_only_due_messages(client, auth_headers, monkeypatch):
    _enable_wechaty(monkeypatch)
    sent: list[tuple[str, str]] = []
    monkeypatch.setattr(
        "sales_fastapi.wechaty.WechatyGateway.send_text",
        lambda self, to, text: sent.append((to, text)) or {"id": "gw-1"},
    )
    _queue(client, auth_headers, when=datetime.utcnow() - timedelta(minutes=5))

    body = client.post("/api/events/messages/send-due", headers=auth_headers).json()
    assert body == {"sent": 3, "failed": 0, "skipped": 0}
    assert len(sent) == 3
    assert all("{{" not in text for _, text in sent)

    listed = client.get("/api/events/messages/all", headers=auth_headers).json()
    assert listed["counts"] == {"sent": 3}
    assert all(row["sent_at"] for row in listed["items"])


def test_send_due_keeps_future_messages_queued(client, auth_headers, monkeypatch):
    _enable_wechaty(monkeypatch)
    monkeypatch.setattr(
        "sales_fastapi.wechaty.WechatyGateway.send_text",
        lambda self, to, text: {"id": "gw-1"},
    )
    _queue(client, auth_headers, when=datetime.utcnow() + timedelta(days=2))
    assert client.post("/api/events/messages/send-due", headers=auth_headers).json()["sent"] == 0
    assert client.get("/api/events/messages/all", headers=auth_headers).json()["counts"] == {
        "queued": 3
    }


def test_send_one_records_gateway_failure(client, auth_headers, monkeypatch):
    _enable_wechaty(monkeypatch)

    def boom(self, to, text):
        raise RuntimeError("gateway down")

    monkeypatch.setattr("sales_fastapi.wechaty.WechatyGateway.send_text", boom)
    _queue(client, auth_headers, when=datetime.utcnow() - timedelta(minutes=1))
    message_id = client.get("/api/events/messages/all", headers=auth_headers).json()["items"][0][
        "id"
    ]

    response = client.post(f"/api/events/messages/{message_id}/send", headers=auth_headers)
    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is False
    assert body["status"] == "failed"
    assert "gateway down" in body["error"]
    assert client.get("/api/events/summary", headers=auth_headers).json()["failed_messages"] == 1


def test_send_one_refuses_a_future_scheduled_message(client, auth_headers):
    _queue(client, auth_headers, when=datetime.utcnow() + timedelta(hours=1))
    message_id = client.get("/api/events/messages/all", headers=auth_headers).json()["items"][0][
        "id"
    ]
    response = client.post(f"/api/events/messages/{message_id}/send", headers=auth_headers)
    assert response.status_code == 409


def test_cancelled_message_cannot_be_sent(client, auth_headers, monkeypatch):
    _enable_wechaty(monkeypatch)
    monkeypatch.setattr(
        "sales_fastapi.wechaty.WechatyGateway.send_text",
        lambda self, to, text: {"id": "gw-1"},
    )
    _queue(client, auth_headers, when=datetime.utcnow() - timedelta(minutes=1))
    queued = client.get("/api/events/messages/all", headers=auth_headers).json()["items"]
    for row in queued:
        assert (
            client.post(
                f"/api/events/messages/{row['id']}/cancel", headers=auth_headers
            ).status_code
            == 200
        )

    assert (
        client.post(
            f"/api/events/messages/{queued[0]['id']}/send", headers=auth_headers
        ).status_code
        == 409
    )
    assert client.post("/api/events/messages/send-due", headers=auth_headers).json()["sent"] == 0
    counts = client.get("/api/events/messages/all", headers=auth_headers).json()["counts"]
    assert counts == {"cancelled": 3}


def test_messages_are_tenant_scoped(client, auth_headers, monkeypatch):
    _enable_wechaty(monkeypatch)
    monkeypatch.setattr(
        "sales_fastapi.wechaty.WechatyGateway.send_text",
        lambda self, to, text: {"id": "gw-1"},
    )
    _queue(client, auth_headers)
    other_token = client.post(
        "/api/auth/google", json={"id_token": "dev:other@kalisoftai.com"}
    ).json()
    other = {"Authorization": f"Bearer {other_token['access_token']}"}
    message_id = client.get("/api/events/messages/all", headers=auth_headers).json()["items"][0][
        "id"
    ]

    assert client.get("/api/events/messages/all", headers=other).json()["items"] == []
    assert (
        client.post(f"/api/events/messages/{message_id}/cancel", headers=other).status_code == 404
    )
    assert client.post(f"/api/events/messages/{message_id}/send", headers=other).status_code == 404
    # The other tenant has nothing queued, so its sweep is a no-op.
    assert client.post("/api/events/messages/send-due", headers=other).json() == {
        "sent": 0,
        "failed": 0,
        "skipped": 0,
    }
    assert client.get("/api/events/messages/all", headers=auth_headers).json()["counts"] == {
        "queued": 3
    }
