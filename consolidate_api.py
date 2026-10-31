"""Rebuild sheet 9: unified FastAPI API catalog (Django endpoints consolidated, Django section removed). Push to GCS."""
import os
import google.auth
from google.cloud import storage
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary_v2.xlsx")
if not os.path.exists(XLSX):
    XLSX = os.path.join(BASE, "Kalika_Dashboard_Summary.xlsx")
GCS_DEST = "sales-metadata/Kalika_Dashboard_Summary.xlsx"

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(bold=True, size=11, color="2E74B5")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = load_workbook(XLSX)
if "9_API_Details_Features" in wb.sheetnames:
    del wb["9_API_Details_Features"]
ws = wb.create_sheet("9_API_Details_Features")


def write_table(r, headers, rows):
    for j, h in enumerate(headers):
        c = ws.cell(row=r, column=1 + j, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.border = THIN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = THIN; c.alignment = WRAP
        r += 1
    return r + 1


r = 1
ws.cell(row=r, column=1, value="CURRENT API — DETAILS, FEATURES & DESCRIPTIONS (UNIFIED FASTAPI SERVICE)").font = TITLE_FONT
r += 1
ws.cell(row=r, column=1, value="Service: 'Kalisoft AI Sales Pipeline' v1.0.0 (FastAPI) — Cloud Run 'kalisoft-sales', asia-south1 — all routes under /api | Legacy web-app endpoints (scraper, extractor, pitch studio, campaign engine, CSV manager, Gmail settings) are CONSOLIDATED into this FastAPI service | Updated 30-Oct-2026").font = Font(italic=True, size=9, color="808080")
r += 2

# ------------------------------------------------ A. overview
ws.cell(row=r, column=1, value="A. SERVICE OVERVIEW & SECURITY").font = SUB_FONT
r += 1
r = write_table(r,
    ["Item", "Detail"],
    [
        ["Service name / version", "Kalisoft AI Sales Pipeline — v1.0.0 (FastAPI)"],
        ["Deployment", "Google Cloud Run: service 'kalisoft-sales', region asia-south1 (cloudbuild.fastapi.yaml); Artifact Registry images; separate sales / image / checks build pipelines"],
        ["API base path", "/api — unified routers; interactive docs at /docs & /redoc (non-production)"],
        ["Auth", "Google OAuth login -> session token (POST /api/auth/google); local signup & password-reset; per-user data isolation via get_current_user"],
        ["Middleware / guardrails", "TrustedHostMiddleware (host allow-list), RequestGuardrailsMiddleware (request safety / PII guardrails), CORS restricted to configured origins; methods GET/POST/PATCH/DELETE"],
        ["Data stores & integrations", "SQL DB (Alembic migrations); GCS bucket kalika_enterprises (contact imports, scrape runs, campaign runs, CSV files); Google Sheets sync; Gmail/IMAP/SMTP; WhatsApp via wechaty-gateway; Gemini/Gemma LLM services"],
        ["Governance", "Audit events, retention purge, usage metering, GCP posture checks, guardrailed LLM task runs"],
        ["Consolidation note", "All web-app modules (scraper, extractor, pitch studio, auto campaign engine, CSV manager, Gmail settings) are part of this single FastAPI service — no separate legacy app."],
    ])

# ------------------------------------------------ B. unified catalog
ws.cell(row=r, column=1, value="B. UNIFIED FASTAPI ENDPOINT CATALOG (current)").font = SUB_FONT
r += 1
r = write_table(r,
    ["Module", "Method", "Endpoint", "Description", "Key Features", "Application / Use"],
    [
        # system
        ["system", "GET", "/api/config", "Public runtime config", "Non-secret client config", "Frontend bootstrap"],
        ["system", "GET", "/api/health", "Health check", "DB-ready flag from lifespan init", "Uptime monitoring / Cloud Run probes"],
        ["system", "GET", "/api/integrations", "Integration status list", "Connected external services", "Integrations dashboard"],
        ["system", "POST", "/api/cron/run", "Trigger cron tasks", "Auth-protected scheduler kick-off", "Nightly jobs, catch-up runs"],
        # auth
        ["auth", "POST", "/api/auth/google", "Login with Google", "Verifies Google ID token, issues session token", "User sign-in"],
        ["auth", "GET", "/api/auth/me", "Current user profile", "Returns authenticated user", "Session check, UI identity"],
        ["auth", "POST", "/api/auth/signup", "Local account signup", "Email/password registration", "Non-Google users"],
        ["auth", "POST", "/api/auth/password-reset", "Password reset flow", "Token-based reset (request/confirm/complete)", "Account recovery"],
        # contacts
        ["contacts", "GET", "/api/contacts", "List contacts", "Filters: domain, intent; per-user", "Contact lists, segmentation"],
        ["contacts", "POST", "/api/contacts", "Create contact", "Validated create (201)", "Manual lead entry"],
        ["contacts", "GET", "/api/contacts/stats", "Contact statistics", "Counts by domain/intent/status", "Dashboard KPI cards"],
        ["contacts", "GET/PATCH/DELETE", "/api/contacts/{id}", "Get / update / delete one contact", "Full CRUD on single record", "Contact detail editing"],
        ["contacts", "POST", "/api/contacts/upload", "Bulk upload contacts", "Multi-file upload (CSV/XLSX)", "Importing marketing lists (Pune / marketing lists)"],
        ["contacts", "POST", "/api/contacts/import/gcs", "Import contacts from GCS", "Optional bucket prefix filter", "Direct sync from kalika_enterprises bucket"],
        ["contacts", "GET", "/api/contacts/diagnostics", "Contacts diagnostics", "Data-quality / pipeline checks", "Ops monitoring"],
        # files
        ["files", "GET", "/api/files/csv", "List CSV files", "Per-user, by category (GCS-backed)", "Shared data files"],
        ["files", "GET", "/api/files/csv/count", "CSV file counts", "Counts for all categories", "Dashboard tiles"],
        ["files", "GET", "/api/files/csv/download/{category}/{filename}", "Download a CSV", "Authenticated download from GCS", "Exports"],
        ["files", "DELETE", "/api/files/csv", "Delete a CSV", "Per-user delete", "File cleanup"],
        # email
        ["email (IMAP/SMTP)", "GET/POST", "/api/email/connections", "List / add mail connections", "Per-user SMTP/IMAP creds (Gmail settings consolidated here)", "Mailbox management"],
        ["email (IMAP/SMTP)", "POST", "/api/email/connections/{id}/test", "Test a connection", "Validates login + deliverability", "Setup verification"],
        ["email (IMAP/SMTP)", "GET", "/api/email/connections/{id}/status", "Connection status poll", "Lightweight status after save/test", "Settings UI polling"],
        ["email (IMAP/SMTP)", "POST", "/api/email/connections/{id}/fetch", "Fetch mail via IMAP", "Filterable fetch (EmailFetchIn)", "Lead reply ingestion"],
        ["email (IMAP/SMTP)", "GET", "/api/email/connections/{id}/inbox", "Read inbox", "limit param (default 15)", "Inbox preview in app"],
        ["email (IMAP/SMTP)", "POST", "/api/email/send", "Send email", "Send via chosen connection_id", "Outreach sending"],
        # google email
        ["email/google (Gmail OAuth)", "GET", "/api/email/google/mailboxes", "List linked Gmail mailboxes", "OAuth-connected accounts", "Multi-mailbox sending"],
        ["email/google (Gmail OAuth)", "GET", "/api/email/google/authorize", "Start Gmail OAuth", "Returns consent URL", "Connect new Gmail"],
        ["email/google (Gmail OAuth)", "GET", "/api/email/google/callback", "OAuth callback", "Stores encrypted tokens", "Completes Gmail link"],
        ["email/google (Gmail OAuth)", "POST", "/api/email/google/mailboxes/{id}/fetch", "Fetch Gmail messages", "Gmail API fetch with filters", "Reply/thread sync"],
        # social
        ["social", "GET/POST", "/api/social/linkedin/profiles | /search", "LinkedIn profiles & search", "Stores found profiles", "Prospect discovery"],
        ["social", "GET/POST", "/api/social/reddit/posts | /search", "Reddit posts & search", "Hiring/intent alerts", "Intent signals"],
        ["social", "POST", "/api/social/youtube/search", "YouTube search", "Channel/video discovery", "Research & prospecting"],
        ["social", "POST/GET", "/api/social/whatsapp/messages", "Queue + list WhatsApp messages", "Outbound queue (201 on create)", "WhatsApp follow-ups"],
        # whatsapp
        ["whatsapp", "GET", "/api/whatsapp/status", "Gateway status", "wechaty-gateway health", "Channel monitoring"],
        ["whatsapp", "GET", "/api/whatsapp/messages", "List WhatsApp messages", "limit param (default 50)", "Message history"],
        ["whatsapp", "POST", "/api/whatsapp/send", "Send WhatsApp message", "Guardrailed send via gateway", "Customer follow-up messaging"],
        ["whatsapp", "POST", "/api/whatsapp/webhook", "Inbound WhatsApp webhook", "Receives incoming messages/events", "Two-way conversations, reply capture"],
        # templates
        ["templates", "GET", "/api/templates/meta", "Template metadata", "Channels/types available", "Template UI options"],
        ["templates", "GET/POST", "/api/templates", "List / create templates", "Channel filter; per-user (email pitch templates consolidated here)", "Message template library"],
        ["templates", "GET/PATCH/DELETE", "/api/templates/{id}", "Load / update / delete template", "Full CRUD", "Template maintenance"],
        ["templates", "POST", "/api/templates/seed", "Seed default templates", "Idempotent seeding", "First-run setup"],
        # campaigns
        ["campaigns", "POST", "/api/campaigns", "Create campaign (201)", "Generates per-recipient messages", "AI outreach campaigns"],
        ["campaigns", "GET", "/api/campaigns | /{id}", "List / get campaign", "CampaignDetail with messages", "Campaign management"],
        ["campaigns", "PATCH", "/api/campaigns/{id}/messages/{mid}", "Edit a campaign message", "Human-in-the-loop editing", "Review before send"],
        ["campaigns", "POST", "/api/campaigns/{id}/approve", "Approve campaign", "Approval gate before sending", "Compliance/quality control"],
        ["campaigns", "POST", "/api/campaigns/{id}/send", "Send approved campaign", "Returns send-result stats", "Bulk outreach execution"],
        ["campaigns", "DELETE", "/api/campaigns/{id}", "Delete campaign (204)", "—", "Cleanup"],
        ["campaigns", "GET", "/api/campaigns/track/open/{id}/{email}", "Open-tracking pixel", "Marks message opened per recipient", "Engagement tracking"],
        # campaign engine (consolidated)
        ["campaigns/engine", "GET", "/api/campaigns/engine/config", "Get auto-campaign config", "Saved config + options (templates, categories)", "Auto-campaign setup"],
        ["campaigns/engine", "POST", "/api/campaigns/engine/config", "Save auto-campaign config", "JSON persist", "Auto-campaign setup"],
        ["campaigns/engine", "POST", "/api/campaigns/engine/run", "Start auto-campaign", "Background run using saved config", "Automated campaigns"],
        ["campaigns/engine", "GET", "/api/campaigns/engine/progress", "Live run progress", "Progress of current/last run", "Dashboard progress bar"],
        ["campaigns/engine", "GET", "/api/campaigns/engine/runs", "List past runs", "Newest first, from GCS", "Run history"],
        ["campaigns/engine", "GET", "/api/campaigns/engine/runs/detail", "Run batch detail", "?path=<blob_path> profiles + messages", "Run audit"],
        # scheduler + kpi
        ["scheduler", "POST/GET", "/api/scheduler/run | /status | /tasks | /clear", "Control & inspect scheduler", "Run now, view tasks, clear queue", "Automation ops"],
        ["kpi", "GET", "/api/kpi/overview", "KPI overview", "Headline funnel metrics", "Dashboard top cards"],
        ["kpi", "GET", "/api/kpi/pipeline", "Pipeline KPIs", "Stage-wise counts", "Pipeline view"],
        ["kpi", "GET", "/api/kpi/outreach", "Outreach KPIs", "Sends/opens/replies", "Outreach performance"],
        ["kpi", "GET", "/api/kpi/costs", "Cost KPIs", "LLM/API spend tracking", "Cost control"],
        ["kpi", "GET", "/api/kpi/website", "Website KPIs", "Site analytics summary", "Marketing insights"],
        # scraper (consolidated)
        ["scraper", "POST", "/api/scraper/search/profile", "Profile search", "People/profile scraping", "Lead research"],
        ["scraper", "POST", "/api/scraper/search/domain", "Domain search", "Company-domain scraping", "Account research"],
        ["scraper", "POST", "/api/scraper/analyze", "Analyze a profile", "AI profile analysis", "Lead qualification"],
        ["scraper", "POST", "/api/scraper/pitch", "Generate pitch from profile", "AI pitch generation", "Instant outreach drafts"],
        ["scraper", "POST/GET/PATCH/DELETE", "/api/scraper/contacts (save | list | update | delete)", "Scraped-contacts CRUD", "Save-to-contacts pipeline + CSV export", "Lead capture"],
        ["scraper", "GET", "/api/scraper/api-status", "Check scraper API status", "External scraper API health", "Ops monitoring"],
        ["scraper/scheduler", "GET/POST", "/api/scraper/keywords (add | toggle | delete)", "Scrape-keyword watchlist", "GCS-config storage; active flags", "Daily keyword scraping config"],
        ["scraper/scheduler", "GET/POST", "/api/scraper/locations (add | toggle | delete)", "Scrape-location watchlist", "GCS-config storage; active flags", "Geo-targeted scraping config"],
        ["scraper/scheduler", "GET", "/api/scraper/runs | /runs/detail", "Scrape run history", "Newest first; ?path= batch detail (profiles + messages)", "Run audit"],
        ["scraper/scheduler", "POST", "/api/scraper/run-now", "Run scrape immediately", "Ad-hoc scrape kick-off", "On-demand harvesting"],
        ["scraper/scheduler", "POST", "/api/scraper/scheduler/run", "Scheduled scrape runner", "Daily keyword x location scraping", "Automated lead harvesting"],
        ["scraper/scheduler", "GET", "/api/scraper/progress", "Live scrape progress", "Current/last run progress", "Dashboard progress bar"],
        ["scraper/scheduler", "GET", "/api/scraper/stats", "Aggregated scrape stats", "Tiles + bar-chart data", "Scraper dashboard"],
        # extractor (consolidated)
        ["extractor", "POST", "/api/extractor/search", "Email/supplier search + CSV upload", "Gemini entity extraction; request/quotation classification", "Supplier discovery from mailboxes"],
        ["extractor", "GET", "/api/extractor/download-csv", "Export supplier data", "CSV export of results", "Reporting/sharing"],
        ["extractor", "POST", "/api/extractor/sync-google-sheet", "Full sync to Google Sheet", "Worksheet sync of extracted data", "Team collaboration"],
        # pitch studio (consolidated)
        ["pitch", "POST", "/api/pitch/create", "Create pitch", "Guided pitch creation", "Quick pitch docs"],
        ["pitch", "GET", "/api/pitch/{id}", "View pitch result", "Rendered pitch output", "Pitch review"],
        ["pitch", "POST", "/api/pitch/generate-subject", "AI subject-line generation", "LLM-generated email subjects", "Email pitch studio"],
        ["pitch", "POST", "/api/pitch/enhance", "AI body enhancement", "Gemma/LLM rewrite of email body", "Email pitch studio"],
        # feedback
        ["feedback", "POST", "/api/feedback", "Submit feedback (201)", "Guardrailed intake", "In-app feedback"],
        ["feedback", "GET", "/api/feedback", "List feedback", "Per-user listing", "Feedback review"],
        ["security", "GET", "/api/security/status", "Security status", "Guardrails/PII status", "Compliance check"],
        # governance
        ["governance", "GET", "/api/governance/policy", "Policy view", "Active governance rules", "Admin review"],
        ["governance", "GET", "/api/governance/gcp", "GCP posture", "GCP config checks", "Cloud compliance"],
        ["governance", "GET", "/api/governance/audit", "Audit events list", "AuditEventOut records", "Audit trail"],
        ["governance", "POST", "/api/governance/retention/purge", "Purge per retention policy", "Data lifecycle enforcement", "Privacy compliance"],
        ["governance", "GET", "/api/governance/usage", "Usage metering", "Per-user usage stats", "Quota/billing insight"],
        # llm
        ["llm", "GET", "/api/llm/tasks", "List LLM tasks", "Allowed task catalog", "AI feature discovery"],
        ["llm", "GET", "/api/llm/plan/{task}", "Get task plan", "Prompt/graph plan for a task", "Transparency of AI runs"],
        ["llm", "POST", "/api/llm/run", "Run an LLM task", "Guardrailed execution (LLMRunOut)", "AI content/analysis generation"],
        # events
        ["events", "GET", "/api/events/summary", "Event summary", "Aggregate counts", "Event dashboard"],
        ["events", "GET", "/api/events | /{id}", "List / event detail", "Search query 'q'", "Event management (e.g., Packvision expo)"],
        ["events", "GET", "/api/events/recipients", "List event recipients", "Searchable", "Invitee lists"],
        ["events", "POST", "/api/events/import", "Import recipients workbook", "Upload XLSX of attendees", "Expo lead import"],
        ["events", "GET", "/api/events/messages/all", "List event messages", "Filter by event_id", "Event comms tracking"],
        ["events", "POST", "/api/events/messages/{id}/send | /cancel", "Send or cancel one event message", "Per-message control", "Scheduled follow-ups"],
        ["events", "POST", "/api/events/messages/send-due", "Send all due messages", "limit param (default 50)", "Batch scheduled sending"],
        ["events", "GET", "/api/events/message-template-default", "Default event template", "—", "Template fallback"],
        ["events", "GET", "/api/events/{id}/message-template", "Per-event template", "Optional recipient preview", "Personalized invites"],
        ["events", "POST", "/api/events/{id}/schedule", "Schedule event messages", "EventScheduleIn payload", "Follow-up automation"],
    ])

# ------------------------------------------------ C. planned
ws.cell(row=r, column=1, value="C. PLANNED — QR FEEDBACK SYSTEM APIs (spec: Kalika_Real_Time_Customer_Feedback_System.md)").font = SUB_FONT
r += 1
r = write_table(r,
    ["Planned API", "Purpose / Description", "Key Features", "Application"],
    [
        ["QR Code Service", "Unique QR entry points (general / delivery / product)", "Secure order refs; source identifiers; no personal data in URLs", "Invoices, challans, labels, brochures"],
        ["Customer Feedback API", "Collect ratings + comments in real time", "Conditional follow-up questions; delivery vs product feedback; duplicate protection", "Post-delivery & product feedback"],
        ["Complaint Ticket API", "Create/route/track tickets", "Open→Assigned→In Progress→Resolved→Closed/Reopened; first-response & resolution timestamps; photo attachments", "Issue resolution workflow"],
        ["ERP Lookup API", "Order/customer/product/delivery lookup from Kalika ERP", "Order/invoice references; live price/stock/delivery estimates", "Auto-identify order on scan"],
        ["WhatsApp Messaging (templates)", "7 ready flows: welcome, delivery, product, issue, catalogue, suggestion, thank-you", "Interactive buttons/lists; consent & opt-out; 1 reminder after 24–48h", "Customer engagement automation"],
        ["Catalogue API", "13 product categories with specs", "Images + enquiry + quotation request", "Mobile catalogue via QR"],
        ["Alerts API", "Low rating→quality; delivery complaint→dispatch; repeats→investigation; unassigned→supervisor", "Configurable alert rules", "Ops escalation"],
        ["Feedback Analytics API", "Dashboard KPIs & exports", "Avg rating, distribution, complaints by product/category, response & resolution times, QR scan vs submit rate", "Management dashboard"],
    ])

for col in range(1, 7):
    ws.column_dimensions[get_column_letter(col)].width = [24, 16, 44, 40, 44, 34][col - 1]

# keep sheet order 1..10
order = ["1_Dashboard_Overview", "2_Contacts_Summary", "3_Call_Status_Followups", "4_Sales_2026_Monthly",
         "5_Top_Customers_Products", "6_ERP_Dispatch_Orders", "7_Product_Rate_List", "8_WhatsApp_Data",
         "9_API_Details_Features", "10_Data_Sources_GCS"]
wb._sheets = [wb[n] for n in order if n in wb.sheetnames]

try:
    wb.save(XLSX)
    print("Excel updated:", XLSX)
except PermissionError:
    XLSX = XLSX.replace(".xlsx", "_new.xlsx")
    wb.save(XLSX)
    print("Locked; saved as:", XLSX)

creds, _ = google.auth.default()
client = storage.Client(credentials=creds)
bucket = client.bucket("kalika_enterprises")
blob = bucket.blob(GCS_DEST)
blob.upload_from_filename(XLSX)
print(f"Pushed to gs://kalika_enterprises/{GCS_DEST} ({os.path.getsize(XLSX)} bytes)")
